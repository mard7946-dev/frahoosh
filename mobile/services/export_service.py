"""خروجی و ورودی فارسی اکسل و PDF برای فراهوش، با متن Unicode واقعی و فونت همراه برنامه."""
from pathlib import Path
from datetime import datetime

def app_dir():
    try:
        from kivy.app import App
        p = Path(App.get_running_app().user_data_dir)
    except Exception:
        p = Path.cwd()
    p.mkdir(parents=True, exist_ok=True)
    return p

def persian_font_path():
    candidates = [
        Path(__file__).resolve().parents[1] / "assets" / "NotoSansArabic-Regular.ttf",
        Path(__file__).resolve().parents[1] / "assets" / "BTitrBd.ttf",
    ]
    for path in candidates:
        if path.is_file():
            return str(path)
    raise FileNotFoundError("فونت فارسی همراه برنامه پیدا نشد.")

def _rtl(value):
    """برای PDF، متن فارسی را برای موتور رسم ReportLab آماده می‌کند."""
    text = str(value or "")
    try:
        from arabic_reshaper import reshape
        from bidi.algorithm import get_display
        return get_display(reshape(text))
    except Exception:
        return text

def export_excel(rows, fields, title="فراهوش"):
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill
    from openpyxl.utils import get_column_letter

    path = app_dir() / f"{title}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "داده‌ها"
    ws.sheet_view.rightToLeft = True

    # متن اکسل عمداً reshape نمی‌شود؛ Excel باید Unicode فارسی واقعی را نمایش دهد.
    for j, key in enumerate(fields, 1):
        cell = ws.cell(1, j, str(key or ""))
        cell.font = Font(name="Noto Sans Arabic", size=11, bold=True)
        cell.alignment = Alignment(horizontal="right", vertical="center")
    ws.row_dimensions[1].height = 24

    for i, row in enumerate(rows or [], 2):
        for j, key in enumerate(fields, 1):
            value = (row or {}).get(key, "")
            cell = ws.cell(i, j, "" if value is None else str(value))
            cell.font = Font(name="Noto Sans Arabic", size=10)
            cell.alignment = Alignment(horizontal="right", vertical="center")

    for j, key in enumerate(fields, 1):
        max_len = len(str(key or ""))
        for row in rows or []:
            max_len = max(max_len, len(str((row or {}).get(key, "") or "")))
        ws.column_dimensions[get_column_letter(j)].width = min(max(max_len + 3, 12), 35)

    ws.freeze_panes = "A2"
    wb.save(path)
    return str(path)

def import_excel(path):
    from openpyxl import load_workbook
    values = list(
        load_workbook(path, read_only=True, data_only=True)
        .active.iter_rows(values_only=True)
    )
    if not values:
        return []
    headers = [str(value).strip() for value in values[0] if value not in (None, "")]
    return [
        {
            headers[i]: row[i]
            for i in range(min(len(headers), len(row)))
            if headers[i] and row[i] is not None
        }
        for row in values[1:]
    ]

def export_pdf(rows, fields, title="گزارش فراهوش"):
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.pdfgen import canvas
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    font_path = persian_font_path()
    font_name = "FrahooshPersian"
    if font_name not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont(font_name, font_path))

    path = app_dir() / f"{title.replace(' ', '_')}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.pdf"
    page_size = landscape(A4)
    c = canvas.Canvas(str(path), pagesize=page_size)
    width, height = page_size

    c.setFont(font_name, 12)
    c.drawCentredString(width / 2, height - 25, _rtl(title))

    count = max(1, len(fields))
    cell_width = (width - 30) / count
    y = height - 45
    c.setFont(font_name, 8)

    def draw_row(row, row_y):
        for index, key in enumerate(fields):
            value = str((row or {}).get(key, "") or "")
            # جلوگیری از بیرون‌زدگی متن در جدول PDF
            value = value[:32]
            x = 15 + index * cell_width
            c.rect(x, row_y - 14, cell_width, 14)
            c.drawCentredString(x + cell_width / 2, row_y - 10, _rtl(value))

    draw_row({key: key for key in fields}, y)
    y -= 14

    for row in rows or []:
        if y < 30:
            c.showPage()
            y = height - 30
            c.setFont(font_name, 8)
        draw_row(row, y)
        y -= 14

    c.save()
    return str(path)
